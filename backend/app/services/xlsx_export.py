"""Выгрузка таблиц в Excel (.xlsx): общий модуль для отчётов и карты ошибок.

Лист = заголовок (необязательные строки), пустая строка, шапка таблицы (жирная, закреплена, с фильтром), строки,
примечания. Любой текст — строкой, а не формулой: значения из ФИО и сценариев не должны исполняться при открытии
в Excel (та же защита, что у CSV в `report_export`). Числа остаются числами. Сеть не используется.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_SHEET_TITLE: Final = 31  # ограничение Excel
MIN_WIDTH: Final = 10
MAX_WIDTH: Final = 60
_BAD_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")
HEADER_FILL: Final = PatternFill("solid", fgColor="E5EDF5")


@dataclass(frozen=True, slots=True)
class Sheet:
    title: str
    columns: list[str]
    rows: Sequence[Sequence[Any]]
    heading: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _sheet_title(raw: str, used: set[str]) -> str:
    base = _BAD_TITLE_CHARS.sub(" ", raw).strip().strip("'") or "Лист"
    title = base[:MAX_SHEET_TITLE]
    counter = 2
    while title.lower() in used:
        suffix = f" ({counter})"
        title = base[: MAX_SHEET_TITLE - len(suffix)] + suffix
        counter += 1
    used.add(title.lower())
    return title


def _put(ws: Any, row: int, column: int, value: Any) -> Any:
    cell = ws.cell(row=row, column=column)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # строка, даже если начинается с «=», «+», «-», «@»
    return cell


def _text_length(value: Any) -> int:
    return max((len(part) for part in str(value).splitlines()), default=0) if value is not None else 0


def render_xlsx(sheets: Sequence[Sheet]) -> bytes:
    if not sheets:
        raise ValueError("Нужен хотя бы один лист")
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for sheet in sheets:
        ws = workbook.create_sheet(_sheet_title(sheet.title, used))
        row = 1
        for index, line in enumerate(sheet.heading):
            cell = _put(ws, row, 1, line)
            cell.font = Font(bold=True, size=14) if index == 0 else Font(color="595959")
            row += 1
        if sheet.heading:
            row += 1  # пустая строка перед таблицей
        header_row = row
        for column, name in enumerate(sheet.columns, start=1):
            cell = _put(ws, row, column, name)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for values in sheet.rows:
            row += 1
            for column, value in enumerate(values, start=1):
                cell = _put(ws, row, column, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        last_data_row = row
        for note in sheet.notes:
            row += 2 if note is sheet.notes[0] else 1
            _put(ws, row, 1, note).font = Font(color="595959")
        for column, name in enumerate(sheet.columns, start=1):
            longest = max([_text_length(name), *(_text_length(r[column - 1]) for r in sheet.rows if column <= len(r))])
            ws.column_dimensions[get_column_letter(column)].width = min(max(longest + 2, MIN_WIDTH), MAX_WIDTH)
        ws.freeze_panes = f"A{header_row + 1}"
        if sheet.columns:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(sheet.columns))}{max(last_data_row, header_row)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
