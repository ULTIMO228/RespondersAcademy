import csv
import io

import pytest
from openpyxl import load_workbook

from app.api.deps import Viewer
from app.db.session import get_sessionmaker
from app.models.report import Report
from app.models.session import Attempt, TeacherOverride, TrainingSession
from app.services.report_export import HEADERS, export_data, render_csv
from app.services.xlsx_export import XLSX_MEDIA_TYPE
from tests.conftest import login_as

REPORT_ID = 'rep-2026-09-16-01-u-005'
GROUP_ID = 'rep-2026-09-16-group'


@pytest.mark.parametrize('extension', ['csv', 'pdf', 'xlsx'])
async def test_export_access_and_format(client, extension):
    url = f'http://test/api/v1/reports/{REPORT_ID}/export.{extension}'
    assert (await client.get(url)).status_code == 401
    await login_as(client, 'student2')
    assert (await client.get(url)).status_code == 403
    await login_as(client, 'student')
    response = await client.get(url)
    assert response.status_code == 200, response.text
    assert 'attachment;' in response.headers['content-disposition']
    assert response.headers['cache-control'] == 'no-store'
    if extension == 'csv':
        assert response.content.startswith(b'\xef\xbb\xbf')
        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        assert rows[0] == HEADERS and len(rows) > 1
        assert all(len(row) == len(HEADERS) for row in rows)
    elif extension == 'xlsx':
        assert response.headers['content-type'] == XLSX_MEDIA_TYPE
        sheets = load_workbook(io.BytesIO(response.content))
        ws = sheets.worksheets[0]
        header_row = next(row for row in ws.iter_rows(values_only=True) if row[0] == HEADERS[0])
        assert list(header_row[: len(HEADERS)]) == HEADERS
        assert ws.max_row > 2
    else:
        assert response.content.startswith(b'%PDF-') and b'/FontFile2' in response.content
    group_url = f'http://test/api/v1/reports/{GROUP_ID}/export.{extension}'
    assert (await client.get(group_url)).status_code == 403
    await login_as(client, 'teacher')
    assert (await client.get(group_url)).status_code == 200
    async with get_sessionmaker()() as db:
        report = await db.get(Report, REPORT_ID)
        session = await db.get(TrainingSession, report.session_id)
        original_teacher = session.teacher_id
        session.teacher_id = 'u-003'
        await db.commit()
    try:
        assert (await client.get(url)).status_code == 403
        await login_as(client, 'admin')
        assert (await client.get(url)).status_code == 200
        assert (await client.get(f'http://test/api/v1/reports/missing/export.{extension}')).status_code == 404
    finally:
        async with get_sessionmaker()() as db:
            session = await db.get(TrainingSession, report.session_id)
            session.teacher_id = original_teacher
            await db.commit()


async def test_csv_attempts_norms_override_and_exam(client):
    async with get_sessionmaker()() as db:
        db.add(TeacherOverride(attempt_id='att-01', teacher_id='u-002', score=55, comment='Проверка экспорта', at='2026-09-21T12:00:00+03:00'))
        attempt = await db.get(Attempt, 'att-01')
        session = await db.get(TrainingSession, attempt.session_id)
        session.format, session.exam = 'exam', {'passThreshold': 70}
        session.plan = {'timeNorms': {'primaryReactionSec': 20, 'fullProcessingSec': 120}}
        await db.flush()
        data = await export_data(db, REPORT_ID, Viewer('u-005', 'student'))
        rows = {r[4]: r for r in data['rows']}
        assert rows['att-01'][11:] == [55, 'не сдал']
        assert rows['att-01'][7] == attempt.primary_reaction_ms - 20000
        assert rows['att-01'][9] == attempt.full_processing_ms - 120000
        session.exam = None
        await db.flush()
        without_threshold = await export_data(db, REPORT_ID, Viewer('u-005', 'student'))
        assert all(r[-1] == '' for r in without_threshold['rows'])
        session.format = 'training'
        await db.flush()
        training = await export_data(db, REPORT_ID, Viewer('u-005', 'student'))
        assert all(r[-1] == '' for r in training['rows'])
        await db.rollback()


def test_csv_quotes_and_formula_text():
    raw = render_csv({'rows': [['=SUM(1,2)', 'Иванов, "Иван"\nстрока', -1000]]})
    rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
    assert rows[1] == ["'=SUM(1,2)", 'Иванов, "Иван"\nстрока', '-1000']


async def test_xlsx_has_the_same_rows_as_csv(client):
    await login_as(client, 'student')
    url = f'http://test/api/v1/reports/{REPORT_ID}/export'
    csv_rows = list(csv.reader(io.StringIO((await client.get(f'{url}.csv')).content.decode('utf-8-sig'))))
    ws = load_workbook(io.BytesIO((await client.get(f'{url}.xlsx')).content)).worksheets[0]
    table = [row for row in ws.iter_rows(values_only=True) if row[0] is not None]
    start = next(i for i, row in enumerate(table) if row[0] == HEADERS[0])
    xlsx_rows = [[('' if v is None else str(v)) for v in row[: len(HEADERS)]] for row in table[start : start + len(csv_rows)]]
    assert xlsx_rows == csv_rows


def test_xlsx_never_turns_text_into_formula():
    from app.services.report_export import render_xlsx

    ws = load_workbook(io.BytesIO(render_xlsx({'id': 'r', 'sessionId': 's', 'generatedAt': 't', 'rows': [['=SUM(1,2)'] + [''] * (len(HEADERS) - 1)], 'insights': []}))).worksheets[0]
    cells = [cell for row in ws.iter_rows() for cell in row if cell.value == '=SUM(1,2)']
    assert cells and all(cell.data_type == 's' for cell in cells)
