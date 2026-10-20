"""Общая выгрузка в Excel (`app.services.xlsx_export`): оформление, защита от формул, имена листов."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.services.xlsx_export import XLSX_MEDIA_TYPE, Sheet, render_xlsx


def _load(content: bytes):
    return load_workbook(io.BytesIO(content))


def test_media_type_and_zip_signature():
    assert XLSX_MEDIA_TYPE == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert render_xlsx([Sheet("Лист", ["А"], [["1"]])])[:2] == b"PK"


def test_heading_header_and_rows_layout():
    sheet = Sheet("Карта", ["Этап", "Попыток"], [["Адрес", 10], ["Службы", 7]], heading=["Карта ошибок", "Группа Г-1"])
    ws = _load(render_xlsx([sheet])).active
    assert [ws.cell(1, 1).value, ws.cell(2, 1).value] == ["Карта ошибок", "Группа Г-1"]
    assert ws.cell(3, 1).value is None  # пустая строка перед таблицей
    assert [ws.cell(4, 1).value, ws.cell(4, 2).value] == ["Этап", "Попыток"]
    assert [ws.cell(5, 1).value, ws.cell(5, 2).value] == ["Адрес", 10]
    assert ws.cell(6, 2).value == 7


def test_numbers_stay_numbers_and_none_is_empty():
    ws = _load(render_xlsx([Sheet("Л", ["a", "b", "c"], [[1, 2.5, None]])])).active
    assert ws.cell(2, 1).value == 1 and ws.cell(2, 1).data_type == "n"
    assert ws.cell(2, 2).value == 2.5
    assert ws.cell(2, 3).value is None


def test_header_is_bold_frozen_and_filtered():
    ws = _load(render_xlsx([Sheet("Л", ["a", "b"], [["x", "y"]], heading=["Заголовок"])])).active
    header_row = 3  # заголовок, пустая строка, шапка
    assert ws.cell(header_row, 1).font.bold is True
    assert ws.freeze_panes == f"A{header_row + 1}"
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith(f"A{header_row}")


@pytest.mark.parametrize("text", ["=SUM(A1:A2)", "+1+1", "-2+3", "@cmd", "  =1+1", "\t=1"])
def test_text_is_never_a_formula(text):
    ws = _load(render_xlsx([Sheet("Л", ["a"], [[text]])])).active
    cell = ws.cell(2, 1)
    assert cell.data_type != "f"
    assert cell.value == text


def test_negative_numbers_are_not_touched():
    ws = _load(render_xlsx([Sheet("Л", ["a"], [[-5]])])).active
    assert ws.cell(2, 1).value == -5


def test_sheet_titles_are_valid_and_unique():
    sheets = [
        Sheet("Журнал: группа [1] / 2?*", ["a"], []),
        Sheet("Журнал: группа [1] / 2?*", ["a"], []),
        Sheet("Очень длинное название листа, которое точно длиннее тридцати одного символа", ["a"], []),
    ]
    names = _load(render_xlsx(sheets)).sheetnames
    assert len(names) == 3 and len(set(names)) == 3
    assert all(len(name) <= 31 for name in names)
    assert not any(ch in name for name in names for ch in "[]:*?/\\")


def test_column_widths_fit_content_with_a_cap():
    ws = _load(render_xlsx([Sheet("Л", ["Коротко", "Длинно"], [["а", "б" * 500]])])).active
    assert ws.column_dimensions["A"].width >= len("Коротко")
    assert ws.column_dimensions["B"].width <= 60


def test_empty_rows_still_make_a_valid_file_with_note():
    ws = _load(render_xlsx([Sheet("Л", ["a", "b"], [], notes=["Нет данных"])])).active
    assert ws.cell(1, 1).value == "a"
    assert any(cell.value == "Нет данных" for row in ws.iter_rows() for cell in row)


def test_notes_follow_the_table():
    ws = _load(render_xlsx([Sheet("Л", ["a"], [["x"]], notes=["Легенда: ✓ редко"])])).active
    values = [row[0].value for row in ws.iter_rows()]
    assert values.index("Легенда: ✓ редко") > values.index("x")


def test_no_sheets_is_an_error():
    with pytest.raises(ValueError):
        render_xlsx([])
