from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from app.services import report_export


def _data(source: str | None) -> dict[str, Any]:
    return {
        "id": "group-report",
        "sessionId": "session-1",
        "generatedAt": "2026-10-20",
        "rows": [],
        "insights": ["Повторите адресные карточки."],
        "insightsSource": source,
    }


def test_xlsx_insights_include_human_readable_source() -> None:
    workbook = load_workbook(io.BytesIO(report_export.render_xlsx(_data("local"))))
    sheet = workbook["Выводы"]

    assert list(sheet.iter_rows(min_row=3, values_only=True)) == [
        ("Вывод", "Источник"),
        ("Повторите адресные карточки.", "ИИ · локально"),
    ]


def test_old_insights_without_source_are_labeled_as_template() -> None:
    workbook = load_workbook(io.BytesIO(report_export.render_xlsx(_data(None))))

    assert workbook["Выводы"]["B4"].value == "шаблон"


def test_pdf_insights_show_source(monkeypatch) -> None:
    captured: list[str] = []
    original = report_export.Paragraph

    def capture(text: str, style: Any):
        captured.append(text)
        return original(text, style)

    monkeypatch.setattr(report_export, "Paragraph", capture)
    rendered = report_export.render_pdf(_data("cloud"))

    assert rendered.startswith(b"%PDF-")
    assert "Источник: ИИ · облако" in captured


def test_csv_export_does_not_add_source_when_it_contains_no_ai_text() -> None:
    # The CSV contains attempt rows only; group insights are exported in PDF/XLSX.
    rendered = report_export.render_csv({"rows": []})

    assert "ИИ" not in rendered.decode("utf-8-sig")
